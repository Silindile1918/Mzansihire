import io
import re
from pypdf import PdfReader
from docx import Document

MAX_CV_MB = 5
MAX_PDF_PAGES = 15

# Skills we look for. You can add your own to this list later.
SKILLS = [
    # Tech
    "Python", "Java", "JavaScript", "TypeScript", "C++", "C#", "SQL", "R programming",
    "HTML", "CSS", "PHP", "React", "Node.js", "Django", "Flask", "Git", "Linux",
    "Machine learning", "Data analysis", "Data visualisation", "Power BI", "Tableau",
    "Excel", "Pandas", "NumPy", "Scikit-learn", "TensorFlow", "AWS", "Azure", "Docker",
    "Cybersecurity", "Networking", "Databases", "Web development", "Mobile development",
    "Software testing", "Agile", "Scrum", "UI/UX", "Figma",
    # Business and finance
    "Accounting", "Bookkeeping", "Pastel", "Sage", "SAP", "Financial reporting",
    "Auditing", "Budgeting", "Payroll", "Tax", "Project management", "Stakeholder management",
    "Report writing", "Administration", "Data entry", "Microsoft Office", "Microsoft Word",
    "PowerPoint", "Research", "Negotiation", "Recruitment", "Human resources",
    # Marketing, sales, service
    "Marketing", "Social media", "SEO", "Content writing", "Copywriting", "Sales",
    "Customer service", "Call centre", "Retail", "Inventory management", "Logistics",
    "Supply chain", "Event management",
    # Creative
    "Graphic design", "Photoshop", "Illustrator", "Canva", "Video editing",
    # Engineering and other
    "AutoCAD", "Civil engineering", "Electrical engineering", "Mechanical engineering",
    "Teaching", "Tutoring", "Public speaking",
    # Soft skills
    "Communication", "Teamwork", "Leadership", "Problem solving", "Time management",
    "Critical thinking", "Presentation skills", "Adaptability", "Creativity",
]


def extract_text(uploaded_file):
    """Reads text from a PDF or DOCX. Returns (text, error_message)."""
    if uploaded_file.size > MAX_CV_MB * 1024 * 1024:
        return None, f"That file is too big. Please use one under {MAX_CV_MB}MB."

    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    try:
        # We check the file's first bytes too, not just its name,
        # so a renamed file can't sneak through.
        if name.endswith(".pdf") and data[:5] == b"%PDF-":
            reader = PdfReader(io.BytesIO(data))
            if len(reader.pages) > MAX_PDF_PAGES:
                return None, f"That PDF has too many pages (max {MAX_PDF_PAGES})."
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif name.endswith(".docx") and data[:2] == b"PK":
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
            text = "\n".join(parts)
        else:
            return None, "Please upload a valid PDF or DOCX file (old .doc files aren't supported)."
    except Exception:
        return None, "We couldn't read that file. It may be damaged or password-protected."

    text = text.strip()
    if len(text) < 50:
        return None, (
            "We couldn't find readable text in that file. "
            "Scanned CVs (photos or images) aren't supported yet."
        )
    return text, None


def detect_skills(text):
    """Returns the skills from our list that appear in the text."""
    found = []
    for skill in SKILLS:
        # Whole-word match, so "Java" does not match inside "JavaScript"
        pattern = r"(?<![A-Za-z0-9])" + re.escape(skill) + r"(?![A-Za-z0-9])"
        if re.search(pattern, text, flags=re.IGNORECASE):
            found.append(skill)
    return found


def merge_skills(existing, found):
    """Adds new skills to the user's existing comma-separated list, without duplicates."""
    current = [s.strip() for s in existing.split(",") if s.strip()]
    already = {s.lower() for s in current}
    for skill in found:
        if skill.lower() not in already:
            current.append(skill)
    return ", ".join(current)
